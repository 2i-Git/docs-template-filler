"""Shared test helpers: build a template .docx and a .xlsx entirely in memory.

Keeping the sample data in code (rather than relying on files in ``input/``)
makes the tests self-contained and hermetic.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest
from docx import Document


def build_template_bytes() -> bytes:
    """A template exercising the tricky cases:

    * a placeholder inside a single run,
    * a placeholder **split across several runs** (the run-aware path),
    * a placeholder inside a table cell,
    * a placeholder with **no matching column** (should be reported as missing).
    """
    doc = Document()

    # Placeholder in a single run.
    doc.add_paragraph("Employee: {{Employee Name}}")

    # Placeholder split across runs — build the runs by hand.
    p = doc.add_paragraph()
    p.add_run("Dear ")
    p.add_run("{{")
    p.add_run("First Name")
    p.add_run("}}")
    p.add_run(",")

    # A number, to check formatting, split awkwardly too.
    p2 = doc.add_paragraph()
    p2.add_run("Salary: ")
    p2.add_run("{{Salary")
    p2.add_run("}}")

    # Case/space tolerance: header is "Employee Name".
    doc.add_paragraph("Confirm: {{  employee   name  }}")

    # Unknown placeholder — must be left untouched and reported.
    doc.add_paragraph("Ref: {{Unknown Field}}")

    # Inside a table.
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "{{Employee Name}}"

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_xlsx_bytes(rows=None, headers=None) -> bytes:
    if headers is None:
        headers = ["Employee Name", "First Name", "Salary"]
    if rows is None:
        rows = [
            ["Alice Smith", "Alice", 50000],
            ["Bob Jones", "Bob", 1234.5],
        ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def template_bytes() -> bytes:
    return build_template_bytes()


@pytest.fixture
def xlsx_bytes() -> bytes:
    return build_xlsx_bytes()


def read_docx_text(data: bytes) -> str:
    """Return all paragraph + table text from a .docx given as bytes."""
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
