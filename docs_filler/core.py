"""Core document-filling logic, shared by the CLI and the web app.

The public entry points are :func:`build_documents` (returns one filled
``.docx`` per spreadsheet row, in memory) and :func:`generate_zip` (the same,
packed into a single ``.zip``). Everything operates on **bytes / file-like
objects**, so nothing needs to touch the disk — the web app can process an
upload entirely in memory.

Placeholder convention: anywhere in the template, ``{{Column Header}}`` is
replaced by that column's value for the current row. Matching ignores
capitalisation and extra spaces.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from typing import Callable, Optional

import openpyxl
from docx import Document

# Placeholders look like {{ Header Name }} — spaces inside the braces are ok.
PLACEHOLDER_RE = re.compile(r"\{\{\s*(.+?)\s*\}\}")


class DocsFillerError(Exception):
    """Raised for user-fixable problems (bad spreadsheet, unknown column, …).

    The CLI turns these into a friendly message; the web app turns them into a
    400 response. Either way the message is safe to show to the user.
    """


def normalise(text: str) -> str:
    """Lower-case and collapse whitespace, for tolerant header matching."""
    return re.sub(r"\s+", " ", str(text).strip()).lower()


# ---------------------------------------------------------------------------
# Spreadsheet reading
# ---------------------------------------------------------------------------

def read_records(data, sheet: Optional[str] = None):
    """Return ``(headers, records)`` from an ``.xlsx`` source.

    ``data`` may be a path or any file-like object (e.g. ``BytesIO``) accepted
    by ``openpyxl.load_workbook``.

    headers: list of original header strings, in column order.
    records: list of dicts mapping normalised header -> raw cell value.
    """
    try:
        wb = openpyxl.load_workbook(data, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises a variety of types on bad input
        raise DocsFillerError(
            "Could not read the spreadsheet. Please upload a valid .xlsx file."
        ) from exc

    try:
        ws = wb[sheet] if sheet else wb.active
    except KeyError as exc:
        raise DocsFillerError(f"Worksheet '{sheet}' was not found in the spreadsheet.") from exc

    rows = [
        list(r) for r in ws.iter_rows(values_only=True)
        if any(c is not None and str(c).strip() != "" for c in r)
    ]
    wb.close()
    if not rows:
        raise DocsFillerError("The spreadsheet is empty.")

    header_row = rows[0]
    headers, index_by_key = [], {}
    for idx, cell in enumerate(header_row):
        if cell is None or str(cell).strip() == "":
            continue
        name = str(cell).strip()
        headers.append(name)
        index_by_key[normalise(name)] = idx

    if not headers:
        raise DocsFillerError(
            "Could not find any column headers in the first row of the spreadsheet."
        )

    records = []
    for row in rows[1:]:
        record = {
            key: (row[idx] if idx < len(row) else None)
            for key, idx in index_by_key.items()
        }
        records.append(record)

    if not records:
        raise DocsFillerError("The spreadsheet has headers but no data rows.")

    return headers, records


# ---------------------------------------------------------------------------
# Value formatting
# ---------------------------------------------------------------------------

def format_value(value) -> str:
    """Turn a spreadsheet cell into a human-friendly string."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if value.is_integer():
            return f"{int(value):,}"
        return f"{round(value, 2):,.2f}"
    return str(value)


# ---------------------------------------------------------------------------
# DOCX replacement (run-aware, so text formatting is preserved)
# ---------------------------------------------------------------------------

def iter_paragraphs(document):
    """Yield every paragraph in the document, including tables/headers/footers."""
    def walk(container):
        for para in container.paragraphs:
            yield para
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        for hf in (section.header, section.footer):
            yield from walk(hf)


def fill_paragraph(paragraph, resolve, missing: set) -> None:
    """Replace all {{...}} in one paragraph while keeping run formatting.

    A placeholder may be split across several runs, so we work on the joined
    text and then write the result back across the original runs.
    """
    runs = paragraph.runs
    if not runs:
        return
    texts = [r.text or "" for r in runs]
    full = "".join(texts)

    matches = list(PLACEHOLDER_RE.finditer(full))
    if not matches:
        return

    bounds, pos = [], 0
    for t in texts:
        bounds.append((pos, pos + len(t)))
        pos += len(t)
    total = pos

    def locate(p: int):
        for i, (s, e) in enumerate(bounds):
            if s <= p < e:
                return i, p - s
        return len(bounds) - 1, bounds[-1][1] - bounds[-1][0]

    # Right-to-left so earlier offsets remain valid as we edit.
    for m in reversed(matches):
        found, replacement = resolve(m.group(1))
        if not found:
            missing.add(m.group(1).strip())
            continue
        si, so = locate(m.start())
        if m.end() >= total:
            ei, eo = len(texts) - 1, len(texts[-1])
        else:
            ei, eo = locate(m.end())
        if si == ei:
            texts[si] = texts[si][:so] + replacement + texts[si][eo:]
        else:
            texts[si] = texts[si][:so] + replacement
            for k in range(si + 1, ei):
                texts[k] = ""
            texts[ei] = texts[ei][eo:]

    for run, new_text in zip(runs, texts):
        run.text = new_text


def fill_document(template, record: dict, missing: set) -> Document:
    """Load the template and fill it for one record.

    ``template`` may be a path or a file-like object (e.g. ``BytesIO``).
    """
    doc = Document(template)

    def resolve(placeholder_name: str):
        key = normalise(placeholder_name)
        if key in record:
            return True, format_value(record[key])
        return False, ""

    for paragraph in iter_paragraphs(doc):
        fill_paragraph(paragraph, resolve, missing)
    return doc


# ---------------------------------------------------------------------------
# Orchestration (shared by CLI and web app)
# ---------------------------------------------------------------------------

def safe_filename(text: str, fallback: str) -> str:
    cleaned = re.sub(r"[^\w\-. ]+", "_", str(text).strip()).strip(". ")
    return cleaned or fallback


def build_documents(
    template_bytes: bytes,
    xlsx_bytes: bytes,
    *,
    name_column: Optional[str] = None,
    sheet: Optional[str] = None,
    base_name: str = "document",
    progress: Optional[Callable[[int, int], None]] = None,
):
    """Fill ``template_bytes`` once per row of ``xlsx_bytes``, in memory.

    Returns ``(documents, missing)`` where ``documents`` is a list of
    ``(filename, docx_bytes)`` and ``missing`` is the sorted list of template
    placeholders that had no matching column. ``progress(done, total)`` is
    called after each document is produced.

    Raises :class:`DocsFillerError` for user-fixable problems.
    """
    headers, records = read_records(BytesIO(xlsx_bytes), sheet)

    # Which column drives each output filename.
    if name_column:
        name_key = normalise(name_column)
        if name_key not in {normalise(h) for h in headers}:
            raise DocsFillerError(
                f"Column '{name_column}' was not found. "
                f"Available columns: {', '.join(headers)}"
            )
    else:
        name_key = normalise(headers[0])

    total = len(records)
    documents: list[tuple[str, bytes]] = []
    all_missing: set = set()
    seen_names: dict = {}

    for i, record in enumerate(records, start=1):
        missing: set = set()
        doc = fill_document(BytesIO(template_bytes), record, missing)
        all_missing |= missing

        base = safe_filename(format_value(record.get(name_key)), f"row_{i}")
        filename = f"{base_name} - {base}"
        count = seen_names.get(filename, 0) + 1
        seen_names[filename] = count
        if count > 1:
            filename = f"{filename} ({count})"

        buffer = BytesIO()
        doc.save(buffer)
        documents.append((f"{filename}.docx", buffer.getvalue()))

        if progress is not None:
            progress(i, total)

    return documents, sorted(all_missing)


def generate_zip(
    template_bytes: bytes,
    xlsx_bytes: bytes,
    *,
    name_column: Optional[str] = None,
    sheet: Optional[str] = None,
    base_name: str = "document",
    progress: Optional[Callable[[int, int], None]] = None,
):
    """Like :func:`build_documents`, but pack the results into a single ``.zip``.

    Returns ``(zip_bytes, missing)``.
    """
    documents, missing = build_documents(
        template_bytes, xlsx_bytes,
        name_column=name_column, sheet=sheet,
        base_name=base_name, progress=progress,
    )

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for filename, data in documents:
            zf.writestr(filename, data)
    return buffer.getvalue(), missing
